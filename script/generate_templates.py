from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation


project_root = Path(__file__).resolve().parent.parent
resources_root = project_root / "FinderTools" / "Resources"
resources_root.mkdir(parents=True, exist_ok=True)

Document().save(resources_root / "blank.docx")

workbook = Workbook()
workbook.save(resources_root / "blank.xlsx")

Presentation().save(resources_root / "blank.pptx")

print("Created blank.docx, blank.xlsx, and blank.pptx in FinderTools/Resources")
